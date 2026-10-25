"""Render a reviewed resume draft as an editable DOCX or a PDF."""
from __future__ import annotations

import io
import re
import os
from pathlib import Path
import unicodedata
from difflib import SequenceMatcher
from dataclasses import dataclass
from html import escape
from typing import Literal


@dataclass(frozen=True)
class Rewrite:
    original: str
    rewritten: str


@dataclass(frozen=True)
class RewriteApplication:
    """Result of applying safe replacements to extracted resume text."""

    text: str
    applied: tuple[Rewrite, ...]
    pending: tuple[Rewrite, ...]


_PUNCTUATION_MAP = str.maketrans({
    "，": ",", "。": ".", "：": ":", "；": ";", "！": "!", "？": "?",
    "（": "(", "）": ")", "【": "[", "】": "]", "“": '"', "”": '"',
    "‘": "'", "’": "'", "、": ",", "－": "-", "—": "-", "–": "-",
})
_LOW_MATCH_SCORE = 60
_FEW_APPLIED_REWRITES = 2


def parse_rewrites(items: list[dict]) -> list[Rewrite]:
    """Validate the small, user-confirmed rewrite payload."""
    if len(items) > 50:
        raise ValueError("一次最多应用 50 条改写")
    rewrites: list[Rewrite] = []
    for item in items:
        if not isinstance(item, dict):
            raise ValueError("每条改写必须是对象")
        original = str(item.get("original", "")).strip()
        rewritten = str(item.get("rewritten", "")).strip()
        if not original or not rewritten:
            raise ValueError("每条改写都必须包含 original 和 rewritten")
        if len(original) > 4000 or len(rewritten) > 4000:
            raise ValueError("单条改写长度不能超过 4000 字符")
        rewrites.append(Rewrite(original=original, rewritten=rewritten))
    return rewrites


def _normalise_char(char: str) -> str:
    """Normalize one character while retaining a useful source span."""
    return unicodedata.normalize("NFKC", char).translate(_PUNCTUATION_MAP)


def _normalised_with_spans(text: str, *, compact: bool = False) -> tuple[str, list[tuple[int, int]]]:
    """Return match text plus offsets back into the original extracted text."""
    chars: list[str] = []
    spans: list[tuple[int, int]] = []
    for index, char in enumerate(text):
        expanded = _normalise_char(char)
        if not expanded:
            continue
        if expanded.isspace():
            if compact:
                continue
            if chars and chars[-1] == " ":
                start, _ = spans[-1]
                spans[-1] = (start, index + 1)
            elif chars:
                chars.append(" ")
                spans.append((index, index + 1))
            continue
        for item in expanded:
            chars.append(item)
            spans.append((index, index + 1))

    left = 0
    right = len(chars)
    while left < right and chars[left] == " ":
        left += 1
    while right > left and chars[right - 1] == " ":
        right -= 1
    return "".join(chars[left:right]), spans[left:right]


def _normalise(text: str, *, compact: bool = False) -> str:
    return _normalised_with_spans(text, compact=compact)[0]


def _normalised_span(text: str, target: str) -> tuple[int, int] | None:
    """Locate target after whitespace/punctuation normalization."""
    for compact in (False, True):
        target_normalised = _normalise(target, compact=compact)
        if not target_normalised:
            continue
        source_normalised, spans = _normalised_with_spans(text, compact=compact)
        position = source_normalised.find(target_normalised)
        if position >= 0:
            start = spans[position][0]
            end = spans[position + len(target_normalised) - 1][1]
            return start, end
    return None


def _line_spans(text: str) -> list[tuple[int, int]]:
    """Return non-empty line content spans for conservative fuzzy matching."""
    spans: list[tuple[int, int]] = []
    offset = 0
    for line in text.splitlines(keepends=True):
        content = line.rstrip("\r\n")
        if content.strip():
            bullet = re.match(r"^\s*(?:[-*•▪◦]|\d+[.)、])\s+", content)
            start = offset + (bullet.end() if bullet else 0)
            end = offset + len(content)
            if start < end:
                spans.append((start, end))
        offset += len(line)
    return spans


def _fuzzy_span(text: str, target: str) -> tuple[int, int] | None:
    """Find one high-confidence, line-level fuzzy candidate.

    Fuzzy matching is intentionally conservative. Ambiguous or weak matches are
    left in the pending section instead of risking a replacement in the wrong
    resume section.
    """
    target_normalised = _normalise(target)
    if len(target_normalised) < 12:
        return None
    candidates: list[tuple[float, int, int]] = []
    for start, end in _line_spans(text):
        candidate = _normalise(text[start:end])
        if not candidate:
            continue
        length_ratio = len(candidate) / len(target_normalised)
        if length_ratio < 0.55 or length_ratio > 1.8:
            continue
        score = SequenceMatcher(None, target_normalised, candidate).ratio()
        if score >= 0.90:
            candidates.append((score, start, end))
    candidates.sort(reverse=True)
    if not candidates:
        return None
    best = candidates[0]
    if len(candidates) > 1 and best[0] - candidates[1][0] < 0.04:
        return None
    return best[1], best[2]


def apply_rewrites_with_report(resume_text: str, rewrites: list[Rewrite]) -> RewriteApplication:
    """Apply safe replacements and retain unmatched suggestions for the output."""
    result = resume_text
    applied: list[Rewrite] = []
    pending: list[Rewrite] = []
    for rewrite in rewrites:
        if rewrite.original in result:
            start = result.find(rewrite.original)
            result = result[:start] + rewrite.rewritten + result[start + len(rewrite.original):]
            applied.append(rewrite)
            continue

        span = _normalised_span(result, rewrite.original)
        if span is None:
            span = _fuzzy_span(result, rewrite.original)
        if span is None:
            pending.append(rewrite)
            continue

        start, end = span
        result = result[:start] + rewrite.rewritten + result[end:]
        applied.append(rewrite)

    return RewriteApplication(result, tuple(applied), tuple(pending))


def apply_rewrites(resume_text: str, rewrites: list[Rewrite]) -> str:
    """Apply safe replacements without blocking on unmatched suggestions."""
    return apply_rewrites_with_report(resume_text, rewrites).text


def _append_pending_section(application: RewriteApplication, match_score: int | None) -> str:
    """Append editable alternatives when a source sentence could not be located."""
    if not application.pending:
        return application.text

    is_low_score = match_score is not None and match_score < _LOW_MATCH_SCORE
    if is_low_score and len(application.applied) < _FEW_APPLIED_REWRITES:
        heading = "低匹配度简历的润色备选"
        intro = "以下内容来自已确认的改写建议，未能安全替换原文，请在 Word 中核对后使用。"
    else:
        heading = "待确认改写建议"
        intro = "以下改写未能安全定位到原始文本，请在 Word 中核对后手动替换。"

    entries = [application.text.rstrip(), "", heading, intro]
    for rewrite in application.pending:
        entries.extend([
            f"- {rewrite.rewritten}",
            f"原文参考：{rewrite.original}",
        ])
    return "\n".join(entries)


def _prepared_text(resume_text: str, rewrites: list[Rewrite], match_score: int | None) -> tuple[str, RewriteApplication]:
    application = apply_rewrites_with_report(resume_text, rewrites)
    return _append_pending_section(application, match_score), application


def _is_heading(line: str) -> bool:
    text = line.strip().rstrip(":：")
    headings = {
        "个人信息", "基本信息", "个人简介", "简介", "工作经历", "工作经验",
        "项目经历", "项目经验", "教育经历", "教育背景", "技能", "专业技能",
        "证书", "荣誉", "其他", "summary", "profile", "experience",
        "work experience", "projects", "project experience", "education",
        "skills", "certifications", "awards", "待确认改写建议", "低匹配度简历的润色备选",
        "代表性实践项目", "实践项目", "项目成果", "荣誉与证书", "获奖情况",
    }
    return text.lower() in headings


def _is_bullet(line: str) -> bool:
    return bool(re.match(r"^\s*(?:[-*•▪◦·□☐☑✓✔]|\d+[.)、])\s*", line))


def _clean_bullet(line: str) -> str:
    return re.sub(r"^\s*(?:[-*•▪◦·□☐☑✓✔]|\d+[.)、])\s*", "", line).strip()


def _lines(text: str) -> list[str]:
    cleaned: list[str] = []
    for line in text.splitlines():
        # PDF extraction can leak CID placeholders and non-printing controls
        # into text. Keep tabs/spaces, but discard malformed glyph markers.
        line = re.sub(r"\(cid:\s*\d+\)", "", line, flags=re.IGNORECASE)
        line = "".join(ch for ch in line if ch == "\t" or not unicodedata.category(ch).startswith("C"))
        line = line.replace("\ufffd", "")
        line = line.strip()
        if line and not re.fullmatch(r"[-*•▪◦·□☐☑✓✔]+", line):
            cleaned.append(line)
    return cleaned


def _section_groups(lines: list[str]) -> list[list[str]]:
    """Group each heading with its following content for sensible pagination."""
    groups: list[list[str]] = []
    current: list[str] = []
    for line in lines:
        if current and _is_heading(line):
            groups.append(current)
            current = []
        current.append(line)
    if current:
        groups.append(current)
    return groups


def _pdf_cjk_font() -> str:
    """Register an embeddable Unicode font, with platform-specific defaults."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    configured = os.getenv("JDSCOUT_CJK_FONT")
    candidates = [Path(configured)] if configured else []
    candidates.extend(Path(path) for path in (
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\msyh.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
    ))
    for path in candidates:
        if not path.is_file():
            continue
        name = "ResumeCJK"
        try:
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(path), subfontIndex=0 if path.suffix.lower() == ".ttc" else 0))
            return name
        except Exception as exc:
            if configured and path == candidates[0]:
                raise RuntimeError(f"无法加载 JDSCOUT_CJK_FONT 字体文件：{path}") from exc
    raise RuntimeError(
        "未找到可嵌入的中文字体。请安装 SimHei/Noto Sans CJK，或设置 JDSCOUT_CJK_FONT 指向 TTF/TTC 字体文件。"
    )


def _mixed_markup(text: str, font_name: str = "ResumeCJK") -> str:
    """Escape text and use one embedded Unicode font for stable CJK spacing."""
    return f'<font name="{font_name}">{escape(text)}</font>'


def _resume_markup(text: str, font_name: str, accent: str) -> str:
    """Emphasize common resume labels while keeping user text XML-safe."""
    labels = r"(研究背景|设计方法|现场测试|现场验证|项目职责|项目成果|核心职责|主要成果|教育背景|工作内容|个人优势)[:：]"
    parts: list[str] = []
    cursor = 0
    for match in re.finditer(labels, text):
        parts.append(_mixed_markup(text[cursor:match.start()], font_name))
        parts.append(f'<font color="{accent}"><b>{_mixed_markup(match.group(0), font_name)}</b></font>')
        cursor = match.end()
    parts.append(_mixed_markup(text[cursor:], font_name))
    return "".join(parts)


def render_docx(
    resume_text: str,
    rewrites: list[Rewrite],
    match_score: int | None = None,
    photo_bytes: bytes | None = None,
) -> bytes:
    """Create a simple, editable Word document from the reviewed resume text."""
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("生成 Word 需要安装 python-docx，请先执行 uv sync") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(37, 42, 52)
    normal.paragraph_format.space_after = Pt(5)
    accent = RGBColor(36, 59, 138)
    for style_name in ("Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = accent
        style.font.bold = True
        style.font.size = Pt(12 if style_name == "Heading 2" else 14)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    prepared_text, _ = _prepared_text(resume_text, rewrites, match_score)
    lines = _lines(prepared_text)
    title = lines[0] if lines else "个人简历"
    if photo_bytes:
        header = document.add_table(rows=1, cols=2)
        header.autofit = False
        header.columns[0].width = Inches(5.45)
        header.columns[1].width = Inches(1.0)
        left, right = header.rows[0].cells
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = left.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if photo_bytes else WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        run.bold = True
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(20 if photo_bytes else 19)
        run.font.color.rgb = accent
        right.paragraphs[0].add_run().add_picture(io.BytesIO(photo_bytes), width=Inches(0.85))
        lines = lines[1:]
    else:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        run.bold = True
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(19)
        run.font.color.rgb = accent
        lines = lines[1:]

    for group in _section_groups(lines):
        for line in group:
            if _is_heading(line):
                paragraph = document.add_paragraph()
                paragraph.style = document.styles["Heading 2"]
                paragraph.paragraph_format.keep_with_next = True
                run = paragraph.add_run(line)
                run.font.name = "Microsoft YaHei"
                run.font.color.rgb = accent
                # Thin blue rule echoes the source PDF's section separators.
                p_pr = paragraph._p.get_or_add_pPr()
                borders = p_pr.find("w:pBdr", p_pr.nsmap)
                if borders is None:
                    from docx.oxml import OxmlElement
                    borders = OxmlElement("w:pBdr")
                    p_pr.append(borders)
                from docx.oxml import OxmlElement
                bottom = OxmlElement("w:bottom")
                bottom.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "single")
                bottom.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sz", "6")
                bottom.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space", "2")
                bottom.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color", "243B8A")
                borders.append(bottom)
            elif _is_bullet(line):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.keep_together = True
                paragraph.add_run(_clean_bullet(line))
            else:
                document.add_paragraph(line)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def render_pdf(
    resume_text: str,
    rewrites: list[Rewrite],
    match_score: int | None = None,
    photo_bytes: bytes | None = None,
) -> bytes:
    """Create a clean, selectable PDF using a Unicode CJK font."""
    try:
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.platypus import Image as PdfImage
        from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("生成 PDF 需要安装 reportlab，请先执行 uv sync") from exc

    font_name = _pdf_cjk_font()
    styles = getSampleStyleSheet()
    accent = "#243B8A"
    body = ParagraphStyle(
        "ResumeBody", parent=styles["BodyText"], fontName=font_name,
        fontSize=10, leading=14.7, spaceAfter=1.7 * mm, textColor="#252A34",
    )
    heading = ParagraphStyle(
        "ResumeHeading", parent=body, fontSize=12.2, leading=16,
        textColor=accent, spaceBefore=3.2 * mm, spaceAfter=1.7 * mm, keepWithNext=True,
    )
    title = ParagraphStyle(
        "ResumeTitle", parent=body, alignment=TA_CENTER, textColor=accent,
        fontSize=20, leading=25, spaceAfter=4.5 * mm,
    )
    subheading = ParagraphStyle(
        "ResumeSubheading", parent=body, fontSize=10.6, leading=15,
        textColor="#26334D", spaceBefore=0.9 * mm, spaceAfter=1.8 * mm,
    )

    story = []
    prepared_text, _ = _prepared_text(resume_text, rewrites, match_score)
    lines = _lines(prepared_text)
    title_line = lines[0] if lines else "个人简历"
    story.append(Paragraph(f"<b>{_mixed_markup(title_line, font_name)}</b>", title))
    lines = lines[1:]
    index = 0
    while index < len(lines):
        line = lines[index]
        if not _is_heading(line):
            story.append(Paragraph(_resume_markup(line, font_name, accent), body))
            index += 1
            continue

        heading_text = line.rstrip(":：")
        heading_cell = Table([[Paragraph(f"<b>{_mixed_markup(heading_text, font_name)}</b>", heading)]], colWidths=[None])
        heading_cell.setStyle(TableStyle([
            ("LINEBELOW", (0, 0), (-1, -1), 0.8, accent),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        story.append(heading_cell)
        index += 1

        # Match the source resume's information block: details on the left,
        # portrait on the right. Other sections flow as readable paragraphs.
        if heading_text in {"基本信息", "个人信息"}:
            details: list[str] = []
            while index < len(lines) and not _is_heading(lines[index]):
                details.append(lines[index])
                index += 1
            detail_flowables = [Paragraph(_resume_markup(item, font_name, accent), body) for item in details]
            if photo_bytes:
                portrait = PdfImage(io.BytesIO(photo_bytes), width=26 * mm, height=31 * mm)
                info = Table([[detail_flowables, portrait]], colWidths=[None, 30 * mm])
                info.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 1 * mm),
                    ("TOPPADDING", (0, 0), (-1, -1), 1.5 * mm),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 1 * mm),
                ]))
                story.append(info)
            else:
                story.extend(detail_flowables)
            continue

        while index < len(lines) and not _is_heading(lines[index]):
            content = lines[index]
            if _is_bullet(content):
                bullet_style = ParagraphStyle("ResumeBullet", parent=body, leftIndent=4 * mm, firstLineIndent=-3 * mm)
                story.append(Paragraph(
                    f'<font name="{font_name}" color="{accent}">-</font>  '
                    f'{_resume_markup(_clean_bullet(content), font_name, accent)}', bullet_style
                ))
            elif re.fullmatch(r"[《〈].+[》〉]", content.strip()):
                story.append(Paragraph(f"<b>{_mixed_markup(content, font_name)}</b>", subheading))
            else:
                story.append(Paragraph(_resume_markup(content, font_name, accent), body))
            index += 1

    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=17 * mm, leftMargin=17 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm, title="Tailored Resume",
    )
    doc.build(story)
    return output.getvalue()


def render_resume_with_report(
    resume_text: str,
    rewrites: list[Rewrite],
    output_format: Literal["docx", "pdf"],
    match_score: int | None = None,
    photo_bytes: bytes | None = None,
) -> tuple[bytes, str, str, RewriteApplication]:
    """Render a resume and return replacement status for the API/UI."""
    prepared_text, application = _prepared_text(resume_text, rewrites, match_score)
    if output_format == "docx":
        # Render the already-prepared text without applying the rewrites twice.
        return (
            render_docx(prepared_text, [], match_score=None, photo_bytes=photo_bytes),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "tailored_resume.docx",
            application,
        )
    if output_format == "pdf":
        return (
            render_pdf(prepared_text, [], match_score=None, photo_bytes=photo_bytes),
            "application/pdf",
            "tailored_resume.pdf",
            application,
        )
    raise ValueError("output_format 只能是 docx 或 pdf")


def render_resume(
    resume_text: str,
    rewrites: list[Rewrite],
    output_format: Literal["docx", "pdf"],
    match_score: int | None = None,
    photo_bytes: bytes | None = None,
) -> tuple[bytes, str, str]:
    """Return bytes, media type, and a safe download filename."""
    content, media_type, filename, _ = render_resume_with_report(
        resume_text, rewrites, output_format, match_score, photo_bytes
    )
    return content, media_type, filename
